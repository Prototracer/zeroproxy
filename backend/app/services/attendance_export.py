"""
Service for exporting course attendance to XLSX and CSV.

Features:
  - Fixed columns: Serial No, Roll No, Name, Session date/time columns (Asia/Kolkata),
    Total Present, Total Sessions, Attendance %.
  - Unified effective attendance resolution (latest correction > log > absent).
  - Cell values: 1 for Present, 0 for Absent.
  - Formula injection defense: prepends "'" to any text starting with =, +, -, @, \t, \r.
  - Optional date range filtering (from_date, to_date).
  - Clean openpyxl XLSX styling with auto-sized columns and CSV RFC 4180 output.
"""

import csv
from datetime import datetime, time, timedelta, timezone
import io
from typing import Any

try:
    from zoneinfo import ZoneInfo
    KOLKATA_TZ = ZoneInfo("Asia/Kolkata")
except Exception:
    KOLKATA_TZ = timezone(timedelta(hours=5, minutes=30))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.attendance_correction import AttendanceCorrection
from app.models.attendance_log import AttendanceLog
from app.models.attendance_session import AttendanceSession
from app.models.class_enrollment import ClassEnrollment
from app.models.course import Course
from app.models.user import User

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def sanitize_formula_cell(value: Any) -> Any:
    """Neutralize spreadsheet formula injection by prefixing risky leading chars with a single quote."""
    if isinstance(value, str) and value:
        if value.startswith(FORMULA_PREFIXES):
            return f"'{value}"
    return value


async def get_course_attendance_export_data(
    db: AsyncSession,
    course: Course,
    tenant_id: int,
    from_date: datetime | None = None,
    to_date: datetime | None = None,
) -> tuple[list[str], list[list[Any]], list[AttendanceSession]]:
    """
    Fetch sessions, enrolled students, attendance logs and corrections,
    resolving unified effective attendance for all sessions in date range.
    """
    # 1. Fetch sessions for this course
    sess_query = select(AttendanceSession).where(
        AttendanceSession.course_id == course.id,
        AttendanceSession.tenant_id == tenant_id,
    )

    if from_date is not None:
        sess_query = sess_query.where(AttendanceSession.started_at >= from_date)
    if to_date is not None:
        sess_query = sess_query.where(AttendanceSession.started_at <= to_date)

    sess_query = sess_query.order_by(AttendanceSession.started_at.asc())
    sess_res = await db.execute(sess_query)
    sessions = list(sess_res.scalars().all())

    # 2. Fetch active enrolled students
    enroll_res = await db.execute(
        select(User)
        .join(ClassEnrollment, ClassEnrollment.student_id == User.id)
        .where(
            ClassEnrollment.course_id == course.id,
            ClassEnrollment.tenant_id == tenant_id,
            ClassEnrollment.active.is_(True),
        )
    )
    students = list(enroll_res.scalars().all())
    student_map = {u.id: u for u in students}

    session_ids = [s.id for s in sessions]

    # 3. Fetch attendance logs for these sessions
    logs_by_session_student: set[tuple[int, int]] = set()
    if session_ids:
        logs_res = await db.execute(
            select(AttendanceLog.session_id, AttendanceLog.student_id).where(
                AttendanceLog.session_id.in_(session_ids),
                AttendanceLog.tenant_id == tenant_id,
            )
        )
        for s_id, st_id in logs_res.all():
            logs_by_session_student.add((s_id, st_id))

    # 4. Fetch attendance corrections for these sessions (ordered by id.asc so latest wins)
    corrections_by_session_student: dict[tuple[int, int], str] = {}
    if session_ids:
        corr_res = await db.execute(
            select(AttendanceCorrection).where(
                AttendanceCorrection.session_id.in_(session_ids),
                AttendanceCorrection.tenant_id == tenant_id,
            ).order_by(AttendanceCorrection.id.asc())
        )
        for c in corr_res.scalars().all():
            corrections_by_session_student[(c.session_id, c.student_id)] = c.status

    # 5. Include any extra students found in logs or corrections who may not be currently active
    extra_student_ids = {
        st_id for _, st_id in logs_by_session_student | set(corrections_by_session_student.keys())
        if st_id not in student_map
    }
    if extra_student_ids:
        extra_users_res = await db.execute(
            select(User).where(
                User.id.in_(extra_student_ids),
                User.tenant_id == tenant_id,
            )
        )
        for u in extra_users_res.scalars().all():
            student_map[u.id] = u

    # Sort students by roll_no, then name
    sorted_students = sorted(
        student_map.values(),
        key=lambda u: (u.roll_no or "", u.name or ""),
    )

    # 6. Build headers
    headers = ["Serial No", "Roll No", "Name"]
    for s in sessions:
        # Convert started_at to Asia/Kolkata timezone
        dt_kolkata = s.started_at.astimezone(KOLKATA_TZ)
        headers.append(dt_kolkata.strftime("%Y-%m-%d %H:%M"))

    headers.extend(["Total Present", "Total Sessions", "Attendance %"])

    # 7. Build data rows
    total_sessions_count = len(sessions)
    rows: list[list[Any]] = []

    for idx, student in enumerate(sorted_students, start=1):
        clean_roll = sanitize_formula_cell(student.roll_no or "")
        clean_name = sanitize_formula_cell(student.name or "")
        row = [idx, clean_roll, clean_name]

        present_in_sessions = 0
        for s in sessions:
            corr_status = corrections_by_session_student.get((s.id, student.id))
            if corr_status is not None:
                is_present = 1 if corr_status == "present" else 0
            elif (s.id, student.id) in logs_by_session_student:
                is_present = 1
            else:
                is_present = 0

            row.append(is_present)
            present_in_sessions += is_present

        pct = (
            round((present_in_sessions / total_sessions_count) * 100, 1)
            if total_sessions_count > 0
            else 100.0
        )
        row.append(present_in_sessions)
        row.append(total_sessions_count)
        row.append(f"{pct}%")
        rows.append(row)

    return headers, rows, sessions


def generate_attendance_csv(headers: list[str], rows: list[list[Any]]) -> str:
    """Generate RFC 4180 compliant CSV string."""
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)
    return output.getvalue()


def generate_attendance_xlsx(
    headers: list[str],
    rows: list[list[Any]],
    course_name: str,
    course_code: str,
) -> bytes:
    """Generate styled XLSX spreadsheet bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Roster"

    # Title block
    title_text = f"{course_name} ({course_code}) — Attendance Report"
    ws.append([sanitize_formula_cell(title_text)])
    ws.append([])  # blank row

    # Headers
    ws.append(headers)

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Format title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 6))
    title_cell = ws.cell(row=1, column=1)
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    title_cell.alignment = left_align

    header_row_idx = 3
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Append data rows
    start_data_row = 4
    for r_idx, r in enumerate(rows, start=start_data_row):
        ws.append(r)
        is_even = (r_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=r_idx, column=col_idx)
            c.border = thin_border
            if is_even:
                c.fill = zebra_fill

            # Formatting per column
            if col_idx in (1, 2):  # S.No, Roll No
                c.alignment = center_align
            elif col_idx == 3:  # Name
                c.alignment = left_align
            else:  # Attendance columns, counts, percentage
                c.alignment = center_align

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < header_row_idx:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
