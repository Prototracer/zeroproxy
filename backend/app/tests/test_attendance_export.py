"""
Unit tests for Attendance Export feature (Phase 4).
Covers:
  - Fixed columns layout in CSV and XLSX
  - Unified effective attendance (1 for Present, 0 for Absent)
  - Formula injection defense (=, +, -, @ prepended with ')
  - Date range filtering (from / to)
  - Validation of format parameter (xlsx / csv)
  - Faculty scoping and ownership check
"""

from datetime import datetime, timezone
import io
from unittest.mock import AsyncMock, MagicMock
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.core.deps import get_current_user
from app.db import get_db
from app.models.attendance_correction import AttendanceCorrection
from app.models.attendance_log import AttendanceLog
from app.models.attendance_session import AttendanceSession
from app.models.course import Course
from app.models.user import User


@pytest.fixture
def faculty_user():
    return User(
        id=20,
        tenant_id=1,
        roll_no="FAC020",
        email="prof.smith@college.edu",
        name="Prof. Smith",
        role="faculty",
        status="active",
        password_hash="hashed",
    )


@pytest.fixture
def other_faculty_user():
    return User(
        id=21,
        tenant_id=1,
        roll_no="FAC021",
        email="prof.jones@college.edu",
        name="Prof. Jones",
        role="faculty",
        status="active",
        password_hash="hashed",
    )


@pytest.fixture
def test_course():
    return Course(
        id=101,
        tenant_id=1,
        faculty_id=20,
        name="Compiler Design",
        code="CS401",
        semester="Semester 7",
        join_code="CD401X",
        join_code_enabled=True,
    )


@pytest.fixture
def student_alice():
    return User(
        id=11,
        tenant_id=1,
        roll_no="25BCS011",
        email="alice@college.edu",
        name="Alice Kumar",
        role="student",
        status="active",
        password_hash="hashed",
    )


@pytest.fixture
def student_bob():
    return User(
        id=12,
        tenant_id=1,
        roll_no="25BCS012",
        email="bob@college.edu",
        name="Bob Sharma",
        role="student",
        status="active",
        password_hash="hashed",
    )


@pytest.fixture
def session_1(test_course):
    return AttendanceSession(
        id=501,
        tenant_id=1,
        course_id=test_course.id,
        faculty_id=20,
        started_at=datetime(2026, 10, 20, 4, 30, tzinfo=timezone.utc),  # 10:00 AM IST
        ended_at=datetime(2026, 10, 20, 5, 30, tzinfo=timezone.utc),
        status="ended",
    )


@pytest.fixture
def session_2(test_course):
    return AttendanceSession(
        id=502,
        tenant_id=1,
        course_id=test_course.id,
        faculty_id=20,
        started_at=datetime(2026, 10, 22, 8, 30, tzinfo=timezone.utc),  # 2:00 PM IST
        ended_at=datetime(2026, 10, 22, 9, 30, tzinfo=timezone.utc),
        status="ended",
    )


def test_export_csv_layout_and_effective_attendance(
    faculty_user, test_course, student_alice, student_bob, session_1, session_2
):
    """
    Export CSV with fixed layout:
    Serial No, Roll No, Name, [sessions in Asia/Kolkata], Total Present, Total Sessions, Attendance %
    """
    app.dependency_overrides[get_current_user] = lambda: faculty_user
    db = MagicMock()
    db.execute = AsyncMock()

    # 1. Course ownership
    mock_course = MagicMock()
    mock_course.scalar_one_or_none.return_value = test_course

    # 2. Sessions query
    mock_sess = MagicMock()
    mock_sess.scalars.return_value.all.return_value = [session_1, session_2]

    # 3. Active students
    mock_st = MagicMock()
    mock_st.scalars.return_value.all.return_value = [student_alice, student_bob]

    # 4. Logs: Alice scanned session 1
    mock_logs = MagicMock()
    mock_logs.all.return_value = [(session_1.id, student_alice.id)]

    # 5. Corrections: Bob corrected to present in session 2
    bob_corr = AttendanceCorrection(
        id=1,
        tenant_id=1,
        session_id=session_2.id,
        student_id=student_bob.id,
        faculty_id=20,
        status="present",
        reason="Medical approved",
        created_at=datetime(2026, 10, 22, 10, 0, tzinfo=timezone.utc),
    )
    mock_corr = MagicMock()
    mock_corr.scalars.return_value.all.return_value = [bob_corr]

    db.execute.side_effect = [
        mock_course,
        mock_sess,
        mock_st,
        mock_logs,
        mock_corr,
    ]

    async def fake_get_db():
        yield db

    app.dependency_overrides[get_db] = fake_get_db

    with TestClient(app) as client:
        res = client.get(f"/faculty/classes/{test_course.id}/export?format=csv")
        assert res.status_code == 200
        assert "text/csv" in res.headers["content-type"]
        assert "attachment" in res.headers["content-disposition"]
        assert f"{test_course.code}_attendance" in res.headers["content-disposition"]

        csv_text = res.text
        lines = [l.strip() for l in csv_text.strip().split("\r\n") if l.strip()]
        assert len(lines) == 3  # header + Alice + Bob

        header = lines[0].split(",")
        assert header[0] == "Serial No"
        assert header[1] == "Roll No"
        assert header[2] == "Name"
        # Session headers in Asia/Kolkata
        assert header[3] == "2026-10-20 10:00"
        assert header[4] == "2026-10-22 14:00"
        assert header[5] == "Total Present"
        assert header[6] == "Total Sessions"
        assert header[7] == "Attendance %"

        # Alice: present in S1 (scan), absent in S2 -> 1, 0, total 1, 2, 50.0%
        alice_row = lines[1].split(",")
        assert alice_row[0] == "1"
        assert alice_row[1] == "25BCS011"
        assert alice_row[2] == "Alice Kumar"
        assert alice_row[3] == "1"
        assert alice_row[4] == "0"
        assert alice_row[5] == "1"
        assert alice_row[6] == "2"
        assert alice_row[7] == "50.0%"

        # Bob: absent in S1, present in S2 (correction) -> 0, 1, total 1, 2, 50.0%
        bob_row = lines[2].split(",")
        assert bob_row[0] == "2"
        assert bob_row[1] == "25BCS012"
        assert bob_row[2] == "Bob Sharma"
        assert bob_row[3] == "0"
        assert bob_row[4] == "1"
        assert bob_row[5] == "1"
        assert bob_row[6] == "2"
        assert bob_row[7] == "50.0%"

    app.dependency_overrides.clear()


def test_export_xlsx(faculty_user, test_course, student_alice, session_1):
    """Export XLSX spreadsheet and parse workbook."""
    app.dependency_overrides[get_current_user] = lambda: faculty_user
    db = MagicMock()
    db.execute = AsyncMock()

    mock_course = MagicMock()
    mock_course.scalar_one_or_none.return_value = test_course

    mock_sess = MagicMock()
    mock_sess.scalars.return_value.all.return_value = [session_1]

    mock_st = MagicMock()
    mock_st.scalars.return_value.all.return_value = [student_alice]

    mock_logs = MagicMock()
    mock_logs.all.return_value = [(session_1.id, student_alice.id)]

    mock_corr = MagicMock()
    mock_corr.scalars.return_value.all.return_value = []

    db.execute.side_effect = [
        mock_course,
        mock_sess,
        mock_st,
        mock_logs,
        mock_corr,
    ]

    async def fake_get_db():
        yield db

    app.dependency_overrides[get_db] = fake_get_db

    with TestClient(app) as client:
        res = client.get(f"/faculty/classes/{test_course.id}/export?format=xlsx")
        assert res.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
        assert ".xlsx" in res.headers["content-disposition"]

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        assert ws.title == "Attendance Roster"

        # Row 1 is Title Block
        assert "Compiler Design" in str(ws.cell(row=1, column=1).value)

        # Row 3 is Header
        headers = [ws.cell(row=3, column=c).value for c in range(1, 8)]
        assert headers[0] == "Serial No"
        assert headers[1] == "Roll No"
        assert headers[2] == "Name"
        assert headers[3] == "2026-10-20 10:00"
        assert headers[4] == "Total Present"
        assert headers[5] == "Total Sessions"
        assert headers[6] == "Attendance %"

        # Row 4 is Student Data
        assert ws.cell(row=4, column=1).value == 1
        assert ws.cell(row=4, column=2).value == "25BCS011"
        assert ws.cell(row=4, column=3).value == "Alice Kumar"
        assert ws.cell(row=4, column=4).value == 1
        assert ws.cell(row=4, column=5).value == 1
        assert ws.cell(row=4, column=6).value == 1
        assert ws.cell(row=4, column=7).value == "100.0%"

    app.dependency_overrides.clear()


def test_export_formula_injection_defense(faculty_user, test_course, session_1):
    """Special formula characters (=, +, -, @) are sanitized with a leading single quote."""
    app.dependency_overrides[get_current_user] = lambda: faculty_user
    db = MagicMock()
    db.execute = AsyncMock()

    malicious_student = User(
        id=99,
        tenant_id=1,
        roll_no="+999MAL",
        email="attacker@college.edu",
        name="=cmd|'/C calc'!A0",
        role="student",
        status="active",
        password_hash="hashed",
    )

    mock_course = MagicMock()
    mock_course.scalar_one_or_none.return_value = test_course

    mock_sess = MagicMock()
    mock_sess.scalars.return_value.all.return_value = [session_1]

    mock_st = MagicMock()
    mock_st.scalars.return_value.all.return_value = [malicious_student]

    mock_logs = MagicMock()
    mock_logs.all.return_value = []

    mock_corr = MagicMock()
    mock_corr.scalars.return_value.all.return_value = []

    db.execute.side_effect = [
        mock_course,
        mock_sess,
        mock_st,
        mock_logs,
        mock_corr,
    ]

    async def fake_get_db():
        yield db

    app.dependency_overrides[get_db] = fake_get_db

    with TestClient(app) as client:
        # Check CSV
        res = client.get(f"/faculty/classes/{test_course.id}/export?format=csv")
        assert res.status_code == 200
        # Sanitized roll_no starts with '+
        assert "'+999MAL" in res.text
        # Sanitized name starts with '=
        assert "'=cmd|'/C calc'!A0" in res.text

    app.dependency_overrides.clear()


def test_export_invalid_format_returns_422(faculty_user, test_course):
    """Unsupported export format returns 422."""
    app.dependency_overrides[get_current_user] = lambda: faculty_user
    db = MagicMock()
    db.execute = AsyncMock()

    mock_course = MagicMock()
    mock_course.scalar_one_or_none.return_value = test_course

    db.execute.side_effect = [mock_course]

    async def fake_get_db():
        yield db

    app.dependency_overrides[get_db] = fake_get_db

    with TestClient(app) as client:
        res = client.get(f"/faculty/classes/{test_course.id}/export?format=pdf")
        assert res.status_code == 422
        assert "Export format must be 'xlsx' or 'csv'" in res.json()["detail"]

    app.dependency_overrides.clear()


def test_export_non_owner_returns_404(other_faculty_user, test_course):
    """Faculty who does not own the course cannot export attendance."""
    app.dependency_overrides[get_current_user] = lambda: other_faculty_user
    db = MagicMock()
    db.execute = AsyncMock()

    mock_course = MagicMock()
    mock_course.scalar_one_or_none.return_value = None  # Scoped lookup fails

    db.execute.side_effect = [mock_course]

    async def fake_get_db():
        yield db

    app.dependency_overrides[get_db] = fake_get_db

    with TestClient(app) as client:
        res = client.get(f"/faculty/classes/{test_course.id}/export?format=xlsx")
        assert res.status_code == 404
        assert "Class not found" in res.json()["detail"]

    app.dependency_overrides.clear()


def test_export_date_range_filtering(faculty_user, test_course, student_alice, session_1):
    """Export attendance with date range filter."""
    app.dependency_overrides[get_current_user] = lambda: faculty_user
    db = MagicMock()
    db.execute = AsyncMock()

    mock_course = MagicMock()
    mock_course.scalar_one_or_none.return_value = test_course

    mock_sess = MagicMock()
    mock_sess.scalars.return_value.all.return_value = [session_1]

    mock_st = MagicMock()
    mock_st.scalars.return_value.all.return_value = [student_alice]

    mock_logs = MagicMock()
    mock_logs.all.return_value = [(session_1.id, student_alice.id)]

    mock_corr = MagicMock()
    mock_corr.scalars.return_value.all.return_value = []

    db.execute.side_effect = [
        mock_course,
        mock_sess,
        mock_st,
        mock_logs,
        mock_corr,
    ]

    async def fake_get_db():
        yield db

    app.dependency_overrides[get_db] = fake_get_db

    with TestClient(app) as client:
        res = client.get(f"/faculty/classes/{test_course.id}/export?from=2026-10-15&to=2026-10-25&format=csv")
        assert res.status_code == 200
        assert "2026-10-20 10:00" in res.text

    app.dependency_overrides.clear()


def test_export_invalid_date_format_returns_422(faculty_user, test_course):
    """Invalid date format returns 422."""
    app.dependency_overrides[get_current_user] = lambda: faculty_user
    db = MagicMock()
    db.execute = AsyncMock()

    mock_course = MagicMock()
    mock_course.scalar_one_or_none.return_value = test_course

    db.execute.side_effect = [mock_course]

    async def fake_get_db():
        yield db

    app.dependency_overrides[get_db] = fake_get_db

    with TestClient(app) as client:
        res = client.get(f"/faculty/classes/{test_course.id}/export?from=invalid-date")
        assert res.status_code == 422
        assert "Invalid date format" in res.json()["detail"]

    app.dependency_overrides.clear()

